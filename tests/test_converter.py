import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('convert_config', ROOT / 'tools/convert_config.py')
converter = importlib.util.module_from_spec(spec)
spec.loader.exec_module(converter)

class ConverterTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.file = Path(self.temp.name) / 'edited.xlsx'
        self.textfile = Path(self.temp.name) / 'texts.xlsx'
        # Freeze adversarial fixtures, so valid designer changes never break tests
        # merely by changing current defaults, row order or spawn group lengths.
        frozen=json.loads((ROOT/'tests/fixtures/default-config-v1.json').read_text())
        self.book=openpyxl.Workbook();self.book.remove(self.book.active)
        self.textbook=openpyxl.Workbook();self.textbook.remove(self.textbook.active)
        for name, headers in converter.HEADERS.items():
            ws=(self.textbook if name=='Texts' else self.book).create_sheet(name)
            ws.append(headers if name=='Documentation' else ['##var',*headers])
            ws.append(['说明'] if name=='Documentation' else ['##',*headers])
        for path,spec in converter.leaves(json.loads((ROOT/'config/schema.json').read_text())):
            if path[0] in ('version','player','rules','map','effects') or path==('themes','floorsPerTheme'):
                value=frozen
                for part in path:value=value[part]
                self.book['Settings'].append([None,'.'.join(path),value])
        for sheet,section in [('Enemies','enemies'),('Items','items')]:
            for key,value in frozen[section].items():
                self.book[sheet].append([None,key,*[value[h] for h in converter.HEADERS[sheet][1:]]])
        for group in frozen['enemySpawns']:
            for item in group['entries']:self.book['EnemySpawns'].append([None,group['fromFloor'],item['id'],item['weight']])
        for item in frozen['itemDrops']:self.book['ItemDrops'].append([None,item['id'],item['weight']])
        for i,key in enumerate(frozen['themes']['order'],1):self.book['Themes'].append([None,i,key])
        for key in ('cave','forest','wetcave','ruins','wooden','modern','cyber','future'):self.book['ThemeCatalog'].append([None,key,f'theme.{key}.name'])
        for item in frozen['localization']['texts']:self.textbook['Texts'].append([None,item['key'],item['zhCN'],item['en']])
        self.addCleanup(self.book.close);self.addCleanup(self.textbook.close)

    def setting(self, key, value):
        for row in self.book['Settings']:
            if row[1].value == key:
                row[2].value = value
                return row[2].coordinate
        self.fail('unknown fixture setting')

    def save(self):
        self.book.save(self.file);self.textbook.save(self.textfile)

    def run_conversion(self):
        self.save()
        return converter.convert(self.file, self.textfile)

    def reject(self, sheet, coordinate):
        with self.assertRaises(converter.ConfigError) as error:
            self.run_conversion()
        self.assertIn(str(self.textfile if sheet=='Texts' else self.file), str(error.exception))
        self.assertIn(f'{sheet}!{coordinate}', str(error.exception))

    def test_edit_flows_through_luban_and_localization(self):
        self.setting('player.hp', 42)
        self.textbook['Texts']['C3']='测试史莱姆'
        data=self.run_conversion()
        self.assertEqual(data['player']['hp'], 42)
        self.assertEqual(data['enemies']['slime']['name'], '测试史莱姆')
        self.assertEqual(data['enemies']['slime']['nameKey'], 'enemy.slime.name')

    def test_invalid_scalar_types(self):
        for value in ['30', True, '=15+15', -1, 3.5, ' 30']:
            with self.subTest(value=value):
                coordinate = self.setting('player.hp', value)
                self.reject('Settings', coordinate)

    def test_formula_in_documentation_is_rejected(self):
        self.book['Documentation']['E2'] = '=1+1'
        self.reject('Documentation', 'E2')

    def test_luban_duplicate_settings(self):
        ws=self.book['Settings'];ws.append([None,'player.hp',30])
        self.reject('Settings', f'B{ws.max_row}')

    def test_luban_duplicate_and_project_missing_ids(self):
        self.book['Enemies']['B4']='slime'
        self.reject('Enemies','B4')
        self.book['Enemies'].delete_rows(4)
        # Remove bat references first so Luban succeeds and project fixed-ID check executes.
        ws=self.book['EnemySpawns']
        for row in range(ws.max_row,2,-1):
            if ws.cell(row,3).value=='bat':ws.delete_rows(row)
        self.reject('Enemies','B1')

    def test_luban_unknown_reference(self):
        self.book['EnemySpawns']['C3']='dragon'
        self.reject('EnemySpawns','C3')

    def test_luban_missing_translation_reference(self):
        self.book['Enemies']['C3']='missing.key'
        self.reject('Enemies','C3')

    def test_weight_group_duplicate(self):
        self.book['EnemySpawns']['C4']='slime'
        self.reject('EnemySpawns','C4')

    def test_zero_weight_group(self):
        self.book['EnemySpawns']['D3']=0;self.book['EnemySpawns']['D4']=0
        self.reject('EnemySpawns','D3')

    def test_every_floor_cell_is_validated_before_grouping(self):
        self.book['EnemySpawns']['B4']=True
        self.reject('EnemySpawns','B4')

    def test_floor_order(self):
        self.book['EnemySpawns']['B3']=2
        self.reject('EnemySpawns','B3')

    def test_pair_and_cap(self):
        self.setting('effects.sleepTurns.min',4)
        coordinate=self.setting('effects.sleepTurns.max',3)
        self.reject('Settings',coordinate)
        self.setting('effects.sleepTurns.max',6);self.setting('player.belly',100)
        coordinate=self.setting('effects.bellyCap',50)
        self.reject('Settings',coordinate)

    def test_structure_and_required_cells(self):
        self.book['Items']['E3']='ignored?';self.reject('Items','E3')
        self.book['Items']['E3']=None;self.book['Items']['C3']=None
        self.reject('Items','C3')

    def test_unknown_sheet(self):
        self.book.create_sheet('Accidental');self.reject('Workbook','A1')

    def test_themes_order(self):
        self.book['Themes']['B3']=True;self.reject('Themes','B3')

    def test_translation_placeholders_and_empty(self):
        ws=self.textbook['Texts'];row=ws.max_row
        ws.cell(row,4).value='Floor {missing}'
        self.reject('Texts',f'D{row}')
        ws.cell(row,4).value='Broken {floor'
        self.reject('Texts',f'D{row}')
        ws.cell(row,4).value=None
        self.reject('Texts',f'D{row}')

    def test_runtime_text_placeholder_contracts(self):
        ws=self.textbook['Texts'];row=ws.max_row
        ws.cell(row,3).value='{missing}';ws.cell(row,4).value='{missing}'
        self.reject('Texts',f'C{row}')
        ws.cell(row,3).value='{seed} {floor}';ws.cell(row,4).value='{seed} {floor}'
        ws['C3']='Slime {name}';ws['D3']='Slime {name}'
        self.reject('Texts','C3')

    def test_current_authoring_books_compile(self):
        converter.convert(ROOT/'config/game.xlsx', ROOT/'config/texts.xlsx')

    def test_extra_translation_key_is_supported(self):
        self.textbook['Texts'].append([None,'future.skill.name','新技能','New skill'])
        result=self.run_conversion()
        self.assertIn('future.skill.name',[x['key'] for x in result['localization']['texts']])

    def test_runtime_payload_size_limit(self):
        for i in range(70):self.textbook['Texts'].append([None,f'large.text{i}','文'*1900,'x'*1900])
        with self.assertRaisesRegex(converter.ConfigError,'256 KiB'):self.run_conversion()

    def test_luban_typed_field_range(self):
        self.book['Enemies']['E3']=-1
        self.reject('Enemies','E3')

    def test_luban_missing_never_falls_back(self):
        self.save()
        with self.assertRaisesRegex(converter.ConfigError,'Luban missing'):
            converter.convert(self.file,self.textfile,luban=Path(self.temp.name)/'missing.dll')

    def test_deterministic_cli_and_failure_preserves_output(self):
        self.save()
        output=Path(self.temp.name)/'game.json'
        command=[sys.executable,str(ROOT/'tools/convert_config.py'),'--workbook',str(self.file),'--texts',str(self.textfile),'--output',str(output)]
        subprocess.run(command,check=True,capture_output=True)
        first=output.read_bytes()
        subprocess.run(command,check=True,capture_output=True)
        self.assertEqual(first,output.read_bytes())
        self.assertEqual(json.loads(first),converter.convert(self.file,self.textfile))
        # Luban writes temporary outputs even on failed references. Never publish them.
        self.book['Enemies']['C3']='missing.key';self.save()
        result=subprocess.run(command,capture_output=True)
        self.assertNotEqual(result.returncode,0)
        self.assertEqual(first,output.read_bytes())

    def test_check_does_not_write(self):
        self.save();output=Path(self.temp.name)/'absent.json'
        subprocess.run([sys.executable,str(ROOT/'tools/convert_config.py'),'--workbook',str(self.file),'--texts',str(self.textfile),'--output',str(output),'--check'],check=True,capture_output=True)
        self.assertFalse(output.exists())

if __name__=='__main__':unittest.main()
